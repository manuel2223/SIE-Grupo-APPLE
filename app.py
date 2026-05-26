import streamlit as st
import pandas as pd
from docxtpl import DocxTemplate
import gspread
from oauth2client.service_account import ServiceAccountCredentials
import io
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
# Importamos la librería para interactuar con los archivos de Google Drive
from googleapiclient.discovery import build

# --- CONFIGURACIÓN DE CREDENCIALES ---
SCOPE = [
    "https://spreadsheets.google.com/feeds",
    "https://www.googleapis.com/auth/drive"
]

def obtener_credenciales():
    info_claves = dict(st.secrets["gcp_service_account"])
    return ServiceAccountCredentials.from_json_keyfile_dict(info_claves, SCOPE)

def conectar_google_sheets():
    creds = obtener_credenciales()
    return gspread.authorize(creds)

# --- FUNCIÓN NUEVA: DESCARGAR PLANTILLA DESDE DRIVE ---
def descargar_plantilla_drive(file_id):
    try:
        creds = obtener_credenciales()
        # Construimos el cliente para hablar con Google Drive
        servicio_drive = build('drive', 'v3', credentials=creds)
        
        # Solicitamos los bytes del archivo .docx
        request = servicio_drive.files().get_media(fileId=file_id)
        file_bytes = io.BytesIO(request.execute())
        return file_bytes
    except Exception as e:
        st.error(f"Error al descargar la plantilla desde Google Drive: {e}")
        return None

# --- ESTILOS DE WORD (Tus funciones de formato idénticas) ---
def aplicar_verdana(paragraph, texto, negrita=False, tamano=10):
    paragraph.clear() 
    run = paragraph.add_run(texto)
    run.font.name = 'Verdana'
    run.font.size = Pt(tamano)
    run.bold = negrita
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Verdana')

def aplicar_borde_grueso(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    right = tcBorders.find(qn('w:right'))
    if right is None:
        right = OxmlElement('w:right')
        tcBorders.append(right)
    right.set(qn('w:val'), 'single')
    right.set(qn('w:sz'), '16') 
    right.set(qn('w:space'), '0')
    right.set(qn('w:color'), '000000')

# --- CARGA DE DATOS EN CACHÉ ---
@st.cache_data
def cargar_datos_nube():
    try:
        client = conectar_google_sheets()
        # ⚠️ REEMPLAZA AQUÍ con tu URL real de Google Sheets
        URL_SHEET = "https://docs.google.com/spreadsheets/d/17TdFQMTpQnGwXEN5z4dUcyB2kLiG9to5l3AbUqJhGrs/"
        spreadsheet = client.open_by_url(URL_SHEET)
        
        df_mun = pd.DataFrame(spreadsheet.worksheet("Municipios").get_all_records())
        df_per = pd.DataFrame(spreadsheet.worksheet("Personas").get_all_records())
        df_car = pd.DataFrame(spreadsheet.worksheet("Cargos").get_all_records())
        
        municipios = df_mun['Municipio'].tolist()
        dicc_atributos = dict(zip(df_mun['Municipio'], df_mun['Atributo']))
        personas = df_per['Nombre'].tolist()
        cargos = df_car['Cargo'].tolist()
        
        return municipios, dicc_atributos, personas, cargos
    except Exception as e:
        st.error(f"Error crítico al conectar con la base de datos de Google: {e}")
        return [], {}, [], []

# --- DISEÑO DE LA INTERFAZ WEB ---
st.set_page_config(page_title="Generador de Oficios Pro", page_icon="📄", layout="centered")

st.title("📄 Generador de Oficios Pro")
st.write("Rellena los datos para generar el documento oficial editable.")

municipios, dicc_atributos, lista_personas, lista_cargos = cargar_datos_nube()

if municipios:
    st.subheader("Datos del Oficio")
    expediente = st.text_input("Nº Expediente", placeholder="Ej: 2026-001")
    municipio_sel = st.selectbox("Municipio / Entidad", [""] + municipios)
    asunto = st.text_input("Asunto", placeholder="Escribe el asunto...").upper()
    
    st.subheader("👥 Asistentes")
    
    if "num_asistentes" not in st.session_state:
        st.session_state.num_asistentes = 1
        
    asistentes_datos = []
    
    for i in range(st.session_state.num_asistentes):
        col1, col2 = st.columns([2, 1])
        with col1:
            p_sel = st.selectbox(f"Persona {i+1}", [""] + lista_personas, key=f"per_{i}")
        with col2:
            c_sel = st.selectbox(f"Cargo {i+1}", [""] + lista_cargos, key=f"car_{i}")
        if p_sel and c_sel:
            asistentes_datos.append({'nombre': p_sel, 'cargo': c_sel})
            
    if st.button("➕ Añadir Persona"):
        st.session_state.num_asistentes += 1
        st.rerun()

    st.markdown("---")
    
    # El botón de descarga ahora solo pide el expediente, municipio, asunto y los asistentes
    if expediente and municipio_sel and asunto and asistentes_datos:
        
        municipio_mayus = municipio_sel.upper()
        tipo_entidad = str(dicc_atributos.get(municipio_sel, "Municipio")).strip().lower()
        
        if tipo_entidad == "municipio":
            municipio_texto_final = f"AYUNTAMIENTO DE {municipio_mayus}"
        elif tipo_entidad == "mancomunidad":
            municipio_texto_final = f"MANCOMUNIDAD DE {municipio_mayus}"
        elif tipo_entidad == "ela":
            municipio_texto_final = f"ENTIDAD LOCAL AUTÓNOMA DE {municipio_mayus}"
        else:
            municipio_texto_final = municipio_mayus

        titulares = []
        suplentes = []
        for asis in asistentes_datos:
            if "suplente" in asis['cargo'].lower():
                suplentes.append(asis)
            else:
                titulares.append(asis)

        try:
            # El ID puro extraído de tu enlace
            ID_PLANTILLA_DRIVE = "1ulMsldNUX0zzt3T0zlsNAmSbENxGkWOm"
            
            # Descargamos el archivo desde Drive en segundo plano
            archivo_memoria = descargar_plantilla_drive(ID_PLANTILLA_DRIVE)
            
            if archivo_memoria:
                doc = DocxTemplate(archivo_memoria)
                context = {
                    'nexpediente': expediente,
                    'municipio': municipio_texto_final,
                    'asunto': asunto
                }
                doc.render(context)

                # Inyección de tablas (Tu algoritmo exacto)
                for table in doc.docx.tables:
                    if len(table.columns) >= 2:
                        tblPr = table._element.xpath('w:tblPr')
                        if tblPr and tblPr[0].xpath('w:tblW'):
                            tblPr[0].xpath('w:tblW')[0].set(qn('w:type'), 'auto')
                            tblPr[0].xpath('w:tblW')[0].set(qn('w:w'), '0')

                        ancho_c1, ancho_c2 = Inches(1.8), Inches(2.0)
                        table.columns[0].width = ancho_c1
                        table.columns[1].width = ancho_c2
                        table.alignment = WD_TABLE_ALIGNMENT.CENTER
                        table.rows[0].cells[0].width = ancho_c1
                        table.rows[0].cells[1].width = ancho_c2

                        table.rows[0].cells[1].text = "" 
                        p_tit = table.rows[0].cells[0].paragraphs[0]
                        p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        aplicar_verdana(p_tit, "TITULARES", negrita=True, tamano=10)

                    for t in titulares:
                        nueva_fila = table.add_row()
                        nueva_fila.cells[0].width = ancho_c1
                        nueva_fila.cells[1].width = ancho_c2
                        aplicar_verdana(nueva_fila.cells[0].paragraphs[0], t['cargo'].upper(), negrita=True, tamano=9)
                        aplicar_verdana(nueva_fila.cells[1].paragraphs[0], t['nombre'], negrita=False, tamano=9)
                        aplicar_borde_grueso(nueva_fila.cells[0])

                    if suplentes:
                        fila_tit_sup = table.add_row()
                        fila_tit_sup.cells[0].width = ancho_c1
                        fila_tit_sup.cells[1].width = ancho_c2
                        fila_tit_sup.cells[1].text = "" 
                        p_sup = fila_tit_sup.cells[0].paragraphs[0]
                        p_sup.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        aplicar_verdana(p_sup, "SUPLENTES", negrita=True, tamano=10)

                        for s in suplentes:
                            fila_sup = table.add_row()
                            fila_sup.cells[0].width = ancho_c1
                            fila_sup.cells[1].width = ancho_c2
                            fila_sup.cells[0].text = "" 
                            aplicar_verdana(fila_sup.cells[1].paragraphs[0], s['nombre'], negrita=False, tamano=9)
                            aplicar_borde_grueso(fila_sup.cells[0])
                    break

                buffer_word = io.BytesIO()
                doc.save(buffer_word)
                buffer_word.seek(0)

                st.success("✨ ¡Oficio procesado con éxito!")
                st.download_button(
                    label="📥 Descargar Oficio Relleno (.docx)",
                    data=buffer_word,
                    file_name=f"Oficio - {municipio_mayus} - Expediente {expediente}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
        except Exception as e:
            st.error(f"Error al renderizar el documento: {e}")
    else:
        st.info("Completa todos los campos superiores y añade al menos un asistente para desbloquear el botón de descarga.")